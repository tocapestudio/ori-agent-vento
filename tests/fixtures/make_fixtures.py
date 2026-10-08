"""Genera archivos de prueba con datos optométricos distintivos."""
from pathlib import Path

import docx
import fitz
import openpyxl
from pptx import Presentation

OUT = Path(__file__).parent


def text_page(pdf, text, fontsize=13):
    page = pdf.new_page(width=595, height=842)
    page.insert_textbox(fitz.Rect(50, 60, 545, 800), text, fontsize=fontsize, fontname="helv")
    return page


def render_png(text):
    tmp = fitz.open()
    text_page(tmp, text, fontsize=18)
    return tmp[0].get_pixmap(dpi=110).tobytes("png")


# 1. PDF con texto
pdf = fitz.open()
text_page(pdf, "Guía clínica de ambliopía - Centro Orión\n\nIntroducción: la ambliopía es una reducción de la agudeza "
               "visual mejor corregida sin causa orgánica aparente.\n\nSe revisa la refracción bajo cicloplejía.")
text_page(pdf, "Protocolo Orion-Alfa\n\nEn el protocolo Orion-Alfa, la oclusión para ambliopía refractiva moderada se "
               "pauta 3 horas diarias durante 14 semanas, con revisión cada 4 semanas.\n\nSi tras 8 semanas la mejora "
               "es inferior a 1 línea, se añade terapia dicóptica.")
pdf.save(OUT / "guia_ambliopia.pdf")

# 2. PDF escaneado (solo imagen)
scan = fitz.open()
png = render_png("INFORME ESCANEADO - Gabinete Optométrico\n\nPaciente Zeta-17 (anonimizado)\n\nEl punto próximo de "
                 "convergencia del paciente Zeta-17 fue de 11 cm, con recobro a 15 cm.\n\nDiagnóstico orientativo: "
                 "insuficiencia de convergencia.")
page = scan.new_page(width=595, height=842)
page.insert_image(page.rect, stream=png)
scan.save(OUT / "informe_escaneado.pdf")

# 3. DOCX
d = docx.Document()
d.add_heading("Protocolo de vergencias", 1)
d.add_paragraph("Este documento describe ejercicios de vergencias para terapia visual en consulta.")
d.add_heading("Ficha Brock Turquesa", 2)
d.add_paragraph("La ficha Brock Turquesa se realiza con una cuerda de 3 metros y 5 cuentas de colores, "
                "durante 6 minutos por sesión.")
d.save(OUT / "protocolo_vergencias.docx")

# 4. XLSX
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Normas"
ws.append(["Prueba", "Valor norma", "Fuente"])
ws.append(["Flexibilidad acomodativa monocular (club Delfín)", "11 cpm", "Normativa interna club Delfín"])
ws.append(["Flexibilidad acomodativa binocular (club Delfín)", "8 cpm", "Normativa interna club Delfín"])
ws2 = wb.create_sheet("Notas")
ws2.append(["Comentario"])
ws2.append(["Usar flipper de +/-2.00 D"])
wb.save(OUT / "normas_acomodacion.xlsx")

# 5. PPTX
prs = Presentation()
for title, body in [
    ("Visión deportiva", "Introducción al entrenamiento visual deportivo."),
    ("Evaluación", "Tiempo de reacción visual y visión periférica."),
    ("Protocolo Halcón", "El protocolo Halcón para porteros de fútbol usa 4 sesiones semanales de 25 minutos "
                         "con luces Fitlight."),
]:
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    s.placeholders[1].text = body
prs.save(OUT / "vision_deportiva.pptx")

# 6. PNG
(OUT / "tabla_estereopsis.png").write_bytes(render_png(
    "Requisitos visuales - Torneo Cóndor\n\nLa estereopsis mínima requerida en el torneo Cóndor de tiro con "
    "arco es de 40 segundos de arco (test Randot)."))

# 7. Documento de perfil A (para test de ámbito)
d2 = docx.Document()
d2.add_heading("Caso privado perfil A", 1)
d2.add_paragraph("Para el paciente Omega se prescribió un prisma base nasal de 7 dioptrías prismáticas en el ojo izquierdo.")
d2.save(OUT / "caso_omega_perfilA.docx")

print("Fixtures creados en", OUT)
