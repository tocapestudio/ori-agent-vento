"""Light self-test: nested accented folders, scanned PDF (OCR), docx, folder awareness, pipeline error. Cleans up after itself."""
import io
import os
import sys
import time

import docx
import fitz
import requests
from PIL import Image, ImageDraw, ImageFont

API = sys.argv[1].rstrip("/") + "/api"
H = {"Authorization": "Bearer " + requests.post(f"{API}/auth/access", json={"code": "Orion9944+"}).json()["token"]}
s = requests.Session()
s.headers.update(H)


def ok(cond, msg):
    print(("OK   " if cond else "FAIL ") + msg)
    if not cond:
        sys.exit(1)


prof = s.post(f"{API}/profiles", json={"name": "Selftest Biblioteca"}).json()
pid = prof["id"]
try:
    mk = lambda name, parent=None: s.post(f"{API}/folders", json={"name": name, "library": "mine", "profile_id": pid, "parent_id": parent}).json()["id"]
    hub = mk("Barcelona Innovation Hub")
    neuro = mk("Neurobiología y rendimiento deportivo", hub)
    deep = mk("La dificultad y el entorno en la motricidad humana", neuro)
    tec = mk("Tecnología y ciencia aplicada al deporte", hub)
    ocr = mk("EN OCR", tec)

    d = docx.Document()
    d.add_heading("Sensores inerciales en el fútbol", 1)
    d.add_paragraph("Los acelerómetros triaxiales permiten cuantificar la carga externa del jugador durante el partido. "
                    "El módulo explica la frecuencia de muestreo de 100 Hz y la validación frente a fotogrametría.")
    b = io.BytesIO()
    d.save(b)

    img = Image.new("RGB", (1240, 700), "white")
    dr = ImageDraw.Draw(img)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 34) if os.path.exists("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf") else ImageFont.load_default()
    for i, line in enumerate(["La variabilidad de la tarea motriz", "aumenta la dificultad percibida.", "El entorno cambiante exige", "adaptación perceptivo-motora."]):
        dr.text((60, 80 + i * 80), line, fill="black", font=font)
    ib = io.BytesIO()
    img.save(ib, format="PNG")
    pdf = fitz.open()
    pdf.new_page(width=620, height=350).insert_image(fitz.Rect(0, 0, 620, 350), stream=ib.getvalue())
    scanned = pdf.tobytes()

    up = lambda name, data, folder: s.post(f"{API}/documents/upload", files={"files": (name, data)},
                                          data={"library": "mine", "profile_id": pid, "folder_id": folder}).json()[0]["id"]
    ids = [up("Sensores inerciales módulo 3.docx", b.getvalue(), tec), up("Dificultad motriz escaneado.pdf", scanned, deep),
           up("Escaneado roto.pdf", b"%PDF-1.4 esto no es un pdf valido", ocr)]

    end = time.time() + 240
    while time.time() < end:
        docs = {x["id"]: x for x in s.get(f"{API}/documents", params={"library": "mine", "profile_id": pid}).json()}
        if all(docs[i]["status"] != "processing" for i in ids):
            break
        time.sleep(3)
    st = [docs[i]["status"] for i in ids]
    print("estados:", st, "| error:", docs[ids[2]].get("error"))
    ok(st[0] == "ready", "docx en carpeta acentuada → listo")
    ok(st[1] == "ready" and docs[ids[1]]["ocr_used"], "PDF escaneado (OCR) en carpeta anidada → listo")
    ok(st[2] == "error" and docs[ids[2]]["error"], "PDF dañado → estado error con mensaje en español")

    status = s.get(f"{API}/library/status", params={"profile_id": pid}).json()
    ok(status["indexed"] == status["chunks"], f"índice sincronizado {status}")

    def ask(q):
        r = s.post(f"{API}/chat", json={"question": q, "profile_id": pid, "scope": "mine"}).json()
        return r["answer"], r["found_in_docs"] and bool(r["doc_sources"]), r["conversation_id"]

    a, found, conv = ask("¿De qué trata el curso de Tecnología y ciencia aplicada al deporte?")
    print("—", a[:300].replace("\n", " "))
    ok(found and "aceler" in a.lower() and "no lo encuentro" not in a.lower(), "pregunta por carpeta/curso → encuentra el docx, con fuentes y sin 'No lo encuentro'")
    a, found, _ = ask("¿Qué dice el documento escaneado sobre la variabilidad de la tarea?")
    print("—", a[:300].replace("\n", " "))
    ok(found and "variab" in a.lower(), "chat encuentra el texto OCR")
    a, found, _ = ask("¿Qué es la queratometría de Javal en ortoqueratología nocturna?")
    print("—", a[:400].replace("\n", " "))
    low = a.lower()
    ok(not found and "clip" not in low and "no tengo acceso" not in low, "sin resultados: sin 'clip' ni 'no tengo acceso'")
    ok("mi biblioteca" in low and ("3" in a or "tres" in low), "sin resultados: indica ámbito y nº de documentos")
    r = s.post(f"{API}/library/reindex", params={"profile_id": pid}).json()
    ok(r["indexed"] == r["chunks"] and r["requeued"] >= 1, f"reindexar biblioteca {r}")
finally:
    for x in s.get(f"{API}/documents", params={"library": "mine", "profile_id": pid}).json():
        s.delete(f"{API}/documents/{x['id']}")
    for c in s.get(f"{API}/conversations", params={"profile_id": pid}).json():
        s.delete(f"{API}/conversations/{c['id']}", params={"profile_id": pid})
    for f in s.get(f"{API}/folders", params={"library": "mine", "profile_id": pid}).json():
        s.delete(f"{API}/folders/{f['id']}")
    s.delete(f"{API}/profiles/{pid}")
print("ALL OK")
