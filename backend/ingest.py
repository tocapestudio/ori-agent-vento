import asyncio
import io
import logging
import re

import docx
import fitz
import openpyxl
from PIL import Image
from pptx import Presentation

from llm import ocr_image

logger = logging.getLogger("ori.ingest")
SUPPORTED = {"pdf", "docx", "xlsx", "pptx", "png", "jpg", "jpeg", "md"}
OCR_CONCURRENCY = asyncio.Semaphore(2)
OCR_TIMEOUT = 180


def to_png(data: bytes, max_side: int = 1600) -> bytes:
    img = Image.open(io.BytesIO(data))
    img.seek(0)
    img = img.convert("RGB")
    img.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


async def _ocr(png: bytes) -> str:
    async with OCR_CONCURRENCY:
        data = await asyncio.to_thread(to_png, png)
        try:
            return await asyncio.wait_for(ocr_image(data), OCR_TIMEOUT)
        except asyncio.TimeoutError:
            raise ValueError("El reconocimiento de texto (OCR) no respondió a tiempo") from None


def _pdf_pages(path):
    with fitz.open(path) as pdf:
        units, ocr_pages = [], []
        for i, page in enumerate(pdf):
            text = page.get_text().strip()
            if len(text) < 40:
                ocr_pages.append(i)
            units.append([f"p. {i + 1}", text])
    return units, ocr_pages


def _render(path, i):
    with fitz.open(path) as pdf:
        return pdf[i].get_pixmap(dpi=150).tobytes("png")


async def _ocr_page(path, i):
    return await _ocr(await asyncio.to_thread(_render, path, i))


async def _pdf(path):
    units, ocr_pages = await asyncio.to_thread(_pdf_pages, path)
    if ocr_pages:
        logger.info("OCR de %d páginas en %s", len(ocr_pages), path)
        results = await asyncio.gather(*[_ocr_page(path, i) for i in ocr_pages], return_exceptions=True)
        errors = [r for r in results if isinstance(r, BaseException)]
        for i, txt in zip(ocr_pages, results):
            if not isinstance(txt, BaseException):
                units[i][1] = txt
        if errors:
            logger.warning("OCR falló en %d/%d páginas de %s: %r", len(errors), len(ocr_pages), path, errors[0])
            if not any(u[1].strip() for u in units):
                e = errors[0]
                raise ValueError(f"No se pudo leer el texto del PDF escaneado: {getattr(e, 'message', None) or e}")
    return [tuple(u) for u in units], bool(ocr_pages)


def _docx(path):
    d = docx.Document(path)
    sections, current = [], []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        if p.style.name.lower().startswith(("heading", "título", "title")) and current:
            sections.append(current)
            current = []
        current.append(p.text.strip())
    if current:
        sections.append(current)
    units = [(f"sección {i + 1}", "\n".join(s)) for i, s in enumerate(sections)]
    for t_i, tbl in enumerate(d.tables):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in tbl.rows]
        units.append((f"tabla {t_i + 1}", "\n".join(rows)))
    return units


def _xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    units = []
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        if not rows:
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        lines = [" | ".join(header)]
        for r in rows[1:]:
            lines.append(" | ".join(f"{h}: {c}" for h, c in zip(header, r) if c is not None))
        for start in range(0, len(lines), 40):
            block = lines[start:start + 40]
            ref = f"hoja '{ws.title}', filas {start + 1}-{start + len(block)}"
            units.append((ref, f"Hoja {ws.title}\n" + "\n".join(block)))
    return units


def _pptx(path):
    prs = Presentation(path)
    units = []
    for i, slide in enumerate(prs.slides):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notas: " + slide.notes_slide.notes_text_frame.text.strip())
        units.append((f"diapositiva {i + 1}", "\n".join(parts)))
    return units


async def extract_units(path: str, ext: str):
    """Returns (list of (ref, text), ocr_used)."""
    if ext == "pdf":
        return await _pdf(path)
    if ext == "docx":
        return await asyncio.to_thread(_docx, path), False
    if ext == "xlsx":
        return await asyncio.to_thread(_xlsx, path), False
    if ext == "md":
        with open(path, encoding="utf-8") as f:
            return [("chuleta", f.read())], False
    if ext == "pptx":
        return await asyncio.to_thread(_pptx, path), False
    with open(path, "rb") as f:
        return [("imagen", await _ocr(f.read()))], True


def chunk_text(text: str, size: int = 900, overlap: int = 150):
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) <= size:
        return [text] if text else []
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            cut = max(text.rfind(". ", start, end), text.rfind("\n", start, end))
            if cut > start + size // 2:
                end = cut + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in chunks if c]
